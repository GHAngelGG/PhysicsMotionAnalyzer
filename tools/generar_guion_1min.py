"""
Genera el guion ULTRA CORTO (1 minuto) para la exposicion de Jose Angel.
Una sola pagina, letra grande, listo para leer de una mirada.
"""
import os
import sys
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

try:
    sys.stdout.reconfigure(encoding="utf-8")
except Exception:
    pass

OUT = (r"C:\Users\josea\OneDrive\Documentos\U\Fisica I\Proyecto Final"
       r"\Guion_Exposicion_JoseAngel_1MIN.docx")

doc = Document()
section = doc.sections[0]
section.top_margin = Cm(1.8)
section.bottom_margin = Cm(1.8)
section.left_margin = Cm(2.0)
section.right_margin = Cm(2.0)

# Colores
AZUL_OSC  = RGBColor(0x0F, 0x17, 0x2A)
AZUL      = RGBColor(0x25, 0x63, 0xEB)
VERDE     = RGBColor(0x10, 0xB9, 0x81)
NARANJA   = RGBColor(0xF5, 0x9E, 0x0B)
PURPURA   = RGBColor(0x8B, 0x5C, 0xF6)
GRIS      = RGBColor(0x64, 0x74, 0x8B)


def t(texto, size=12, bold=False, italic=False, color=None,
      align=WD_ALIGN_PARAGRAPH.LEFT, space_after=6, indent=0):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_after = Pt(space_after)
    if indent:
        p.paragraph_format.left_indent = Cm(indent)
    r = p.add_run(texto)
    r.font.name = "Calibri"
    r.font.size = Pt(size)
    r.bold = bold
    r.italic = italic
    if color is not None:
        r.font.color.rgb = color
    return p


def fondo(p, hex_color):
    pPr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    pPr.append(shd)


# ─── TITULO ───
t("MI PARTE  ·  1 MINUTO", size=11, bold=True, color=GRIS,
  align=WD_ALIGN_PARAGRAPH.CENTER, space_after=2)
t("EXPLICACIÓN DEL CÓDIGO", size=22, bold=True, color=AZUL_OSC,
  align=WD_ALIGN_PARAGRAPH.CENTER, space_after=14)


# ─── INTRO (5 seg) ───
p = t("« Nuestro sistema tiene 3 partes en el código: una para "
      "encontrar la pelota, una para medirla y una para clasificar "
      "el movimiento. »",
      size=13, bold=True, color=AZUL_OSC, space_after=14)
fondo(p, "DBEAFE")


# ─── CAPTURA 1 (~15 seg) ───
t("🟠  CAPTURA 1  —  Detección por color", size=14, bold=True,
  color=NARANJA, space_after=4)

t("« Esta primera parte se encarga de ENCONTRAR la pelota en el video. »",
  size=13, color=AZUL_OSC, space_after=4, indent=0.5)

t("« El sistema convierte la imagen a un formato llamado HSV, que "
  "separa el color del brillo, y crea una silueta blanca donde está "
  "la pelota y negra en el fondo. Después limpia el ruido para que "
  "quede una silueta clara. »",
  size=12, color=AZUL_OSC, space_after=14, indent=0.5)


# ─── CAPTURA 2 (~15 seg) ───
t("📐  CAPTURA 2  —  Cálculos de física", size=14, bold=True,
  color=VERDE, space_after=4)

t("« Acá CALCULAMOS la velocidad y la aceleración. »",
  size=13, color=AZUL_OSC, space_after=4, indent=0.5)

t("« El sistema toma las posiciones que registró y aplica una fórmula "
  "simple: resta la posición anterior de la siguiente, y la divide entre "
  "el tiempo que pasó. Eso da la velocidad. Y para la aceleración hace lo "
  "mismo, pero con las velocidades. »",
  size=12, color=AZUL_OSC, space_after=14, indent=0.5)


# ─── CAPTURA 3 (~15 seg) ───
t("🎯  CAPTURA 3  —  Clasificación", size=14, bold=True,
  color=PURPURA, space_after=4)

t("« Y esta tercera parte DECIDE automáticamente qué tipo de "
  "movimiento es. »", size=13, color=AZUL_OSC, space_after=4, indent=0.5)

t("« Con la aceleración promedio, el programa se hace 3 preguntas: "
  "si es casi cero, es MRU. Si es cercana a 9.8 — la gravedad — es "
  "Caída Libre. Y cualquier otra aceleración constante es MRUV. »",
  size=12, color=AZUL_OSC, space_after=14, indent=0.5)


# ─── CIERRE (~10 seg) ───
p = t("« En resumen: el código encuentra, mide y clasifica. "
      "Todo automáticamente, cuadro por cuadro, en tiempo real. »",
      size=13, bold=True, color=AZUL_OSC, space_after=14,
      align=WD_ALIGN_PARAGRAPH.CENTER)
fondo(p, "D1FAE5")


# ─── LINEA SEPARADORA ───
p = doc.add_paragraph()
pPr = p._p.get_or_add_pPr()
pbdr = OxmlElement("w:pBdr")
bottom = OxmlElement("w:bottom")
bottom.set(qn("w:val"), "single")
bottom.set(qn("w:sz"), "6")
bottom.set(qn("w:color"), "CBD5E1")
pbdr.append(bottom)
pPr.append(pbdr)


# ─── TARJETA RECORDATORIO ───
t("📝  CHECKLIST RÁPIDO ANTES DE EMPEZAR", size=10, bold=True,
  color=GRIS, space_after=4)

checklist = [
    "✓  La pantalla está en la slide del código",
    "✓  Respiro profundo, hablo despacio",
    "✓  Señalo cada captura cuando la nombro",
    "✓  Termino con la frase verde de cierre",
]
for item in checklist:
    t(item, size=11, color=AZUL_OSC, space_after=2, indent=0.5)

doc.add_paragraph()


# ─── VERSION ULTRA CORTA (back-up) ───
p = t("🚨  SI TE TRABÁS  —  Decí solo esto:",
      size=11, bold=True, color=RGBColor(0xEF, 0x44, 0x44), space_after=4)

p = t("« El código tiene 3 partes: la primera ENCUENTRA la pelota "
      "por su color, la segunda MIDE la velocidad y aceleración, "
      "y la tercera DECIDE si es MRU, MRUV o Caída Libre. »",
      size=12, italic=True, color=AZUL_OSC, space_after=4, indent=0.5)


# Guardar
doc.save(OUT)
print(f"OK: {OUT}")
print(f"Tamaño: {os.path.getsize(OUT)/1024:.1f} KB")
